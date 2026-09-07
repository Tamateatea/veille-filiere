# -*- coding: utf-8 -*-
"""Integre les chaines YouTube officielles des marques (proposition du 06/09).

Entree : le dernier `recherche/comptes_marques_<date>.csv` — une ligne par
marque ou groupe, avec la chaine trouvee et sa certitude (« sur »,
« a verifier », « introuvable »).

Pour chaque chaine retenue, deux effets, chacun a sa place :

  1. La chaine entre dans la table `comptes` de la base, marquee VITRINE
     de son entite (colonne `entite_vitrine`). Le facteur la visite chaque
     nuit ; ce qu'elle publie part dans le flux DECOUVERTE (on y cherche
     des noms de createurs), jamais en detection — c'est la marque qui
     parle, pas un createur.
  2. Le pseudo « @marque » entre dans DICTIONNAIRE.xlsx, type « compte de
     marque », statut `confirme`, force `fort` : une description de
     createur qui mentionne « @danette » est une prise du canal marque par
     @compte (le seul canal marque qui a survecu a la mesure du 06/09, cf.
     Nico d'Estais x @nestleenfrance). Les noms de marque en texte libre
     restent hors detection (decision D-marques).

Par defaut le script SIMULE : il liste ce qu'il ferait, rien n'est ecrit
(critere 11 : l'outil propose, Vincent confirme). Avec `--appliquer`, il
ecrit ; avec `--inclure-a-verifier`, il prend aussi les chaines dont
l'officialite n'est pas prouvee (deconseille sans verification humaine).

Usage :
  python outils/integrer_comptes_marques.py                  # simulation
  python outils/integrer_comptes_marques.py --appliquer      # les « sur »
"""

import argparse
import csv
import datetime as dt
import re
import sqlite3
import sys
import unicodedata
from pathlib import Path

import openpyxl

RACINE = Path(__file__).resolve().parent.parent
CHEMIN_DICO = RACINE / "DICTIONNAIRE.xlsx"
CHEMIN_BASE = RACINE / "donnees" / "veille.sqlite"
DOSSIER_RECHERCHE = RACINE / "recherche"


def normaliser(texte):
    t = unicodedata.normalize("NFD", str(texte or ""))
    t = "".join(c for c in t if unicodedata.category(c) not in ("Mn", "Cf"))
    return " ".join(t.lower().split())


def nombre_abonnes(texte):
    """« 57,5 k abonnés » -> 57500 ; « 502 abonnés » -> 502 ; vide -> None."""
    m = re.match(r"\s*([\d\s,.]+)\s*(k|M)?", str(texte or ""))
    if not m or not m.group(1).strip():
        return None
    valeur = float(m.group(1).replace(" ", "").replace(",", "."))
    facteur = {"k": 1_000, "M": 1_000_000}.get(m.group(2), 1)
    return int(valeur * facteur)


def dernier_fichier():
    fichiers = sorted(DOSSIER_RECHERCHE.glob("comptes_marques_*.csv"))
    if not fichiers:
        sys.exit("aucun recherche/comptes_marques_*.csv")
    return fichiers[-1]


def entites_du_dictionnaire():
    wb = openpyxl.load_workbook(CHEMIN_DICO, read_only=True, data_only=True)
    lignes = list(wb["Entites"].iter_rows(values_only=True))
    signaux = list(wb["Signaux"].iter_rows(values_only=True))
    wb.close()
    entites = {normaliser(l[0]): l[0] for l in lignes[1:] if l[0]}
    entetes = [str(c or "").strip() for c in signaux[0]]
    i_texte = entetes.index("texte")
    textes = {normaliser(l[i_texte]) for l in signaux[1:] if l[i_texte]}
    return entites, textes


def principal():
    parseur = argparse.ArgumentParser()
    parseur.add_argument("--appliquer", action="store_true",
                         help="ecrit dans la base et le dictionnaire "
                              "(sinon : simulation)")
    parseur.add_argument("--inclure-a-verifier", action="store_true",
                         help="prend aussi les chaines « a verifier »")
    args = parseur.parse_args()

    chemin = dernier_fichier()
    with open(chemin, encoding="utf-8-sig") as f:
        lignes = list(csv.DictReader(f))
    certitudes = {"sur"} | ({"a verifier"} if args.inclure_a_verifier else set())
    retenues = [l for l in lignes if l["certitude"] in certitudes
                and l["chaine_trouvee"]]

    entites, textes_dico = entites_du_dictionnaire()
    base = sqlite3.connect(CHEMIN_BASE)
    colonnes = [c[1] for c in base.execute("PRAGMA table_info(comptes)")]
    if "entite_vitrine" not in colonnes and args.appliquer:
        with base:
            base.execute("ALTER TABLE comptes ADD COLUMN entite_vitrine TEXT")
    maintenant = dt.datetime.now().isoformat(timespec="seconds")
    aujourd_hui = dt.date.today().isoformat()

    actions = []  # (marque, chaine, titre, pseudo, action compte, action dico)
    n_comptes_nouveaux = n_comptes_marques = n_pseudos = 0
    inconnues = []
    for l in retenues:
        marque = l["marque"].strip()
        if normaliser(marque) not in entites:
            inconnues.append(marque)
            continue
        entite = entites[normaliser(marque)]
        cid = l["chaine_trouvee"].strip()
        titre = l["titre_de_la_chaine"].strip()
        pseudos = [p.strip() for p in l["pseudo_teste"].split(",") if p.strip()]
        pseudo = pseudos[0] if len(pseudos) == 1 else ""
        existe = base.execute(
            "SELECT nom, abonnes FROM comptes WHERE compte_id = ?",
            (cid,)).fetchone()
        if existe:
            action_compte = f"deja en base ({existe[0]}) -> marquee vitrine"
            n_comptes_marques += 1
        else:
            action_compte = "AJOUTEE a la surveillance, marquee vitrine"
            n_comptes_nouveaux += 1
            n_comptes_marques += 1
        if pseudo and normaliser(pseudo) not in textes_dico:
            action_dico = f"{pseudo} -> compte de marque, confirme, fort"
            n_pseudos += 1
        elif pseudo:
            action_dico = f"{pseudo} deja au dictionnaire"
        else:
            action_dico = "pseudo non resolu (plusieurs testes) : rien"
        actions.append((entite, cid, titre, pseudo, action_compte, action_dico))

        if args.appliquer:
            with base:
                if existe:
                    base.execute(
                        "UPDATE comptes SET entite_vitrine = ?, "
                        "abonnes = COALESCE(NULLIF(abonnes, 0), ?) "
                        "WHERE compte_id = ?",
                        (entite, nombre_abonnes(l["abonnes"]), cid))
                else:
                    base.execute(
                        "INSERT INTO comptes (compte_id, nom, abonnes, "
                        "surveille, ajoute_le, entite_vitrine) "
                        "VALUES (?, ?, ?, 1, ?, ?)",
                        (cid, titre, nombre_abonnes(l["abonnes"]),
                         maintenant, entite))
    base.close()

    if args.appliquer and n_pseudos:
        wb = openpyxl.load_workbook(CHEMIN_DICO)
        ws = wb["Signaux"]
        entetes = [str(c.value or "").strip() for c in ws[1]]
        for entite, cid, titre, pseudo, _, action_dico in actions:
            if not action_dico.startswith(pseudo + " ->"):
                continue
            ligne = {c: "" for c in entetes}
            ligne.update({
                "texte": pseudo, "type_signal": "compte de marque",
                "entite": entite, "plateforme": "YouTube",
                "statut": "confirme", "force": "fort",
                "ajoute_le": aujourd_hui,
                "source": f"{chemin.name}, confirme par Vincent le "
                          f"{aujourd_hui}",
                "commentaire": f"chaine officielle {titre} "
                               f"(youtube.com/channel/{cid})",
            })
            ws.append([ligne.get(c, "") for c in entetes])
        wb.save(CHEMIN_DICO)

    mode = "APPLICATION" if args.appliquer else "SIMULATION (rien n'est ecrit)"
    chemin_md = DOSSIER_RECHERCHE / (
        f"integration_comptes_marques_{aujourd_hui}"
        f"{'' if args.appliquer else '_simulation'}.md")
    with open(chemin_md, "w", encoding="utf-8") as f:
        f.write(f"# Integration des chaines officielles de marques — "
                f"{aujourd_hui} — {mode}\n\n")
        f.write(f"Produit par `outils/integrer_comptes_marques.py` depuis "
                f"`{chemin.name}` (certitudes retenues : "
                f"{', '.join(sorted(certitudes))}).\n\n")
        f.write("| | |\n|---|---:|\n")
        f.write(f"| Chaines retenues | {len(retenues)} |\n")
        f.write(f"| Comptes nouveaux a surveiller | {n_comptes_nouveaux} |\n")
        f.write(f"| Comptes marques vitrine (dont deja en base) | "
                f"{n_comptes_marques} |\n")
        f.write(f"| Pseudos @marque ajoutes au dictionnaire | {n_pseudos} |\n")
        f.write(f"| Marques inconnues du dictionnaire (ignorees) | "
                f"{len(inconnues)} |\n\n")
        if inconnues:
            f.write("Inconnues : " + ", ".join(inconnues) + "\n\n")
        f.write("| Entite | Chaine | Titre | Pseudo | Compte | Dictionnaire |\n")
        f.write("|---|---|---|---|---|---|\n")
        for a in actions:
            f.write("| " + " | ".join(str(x) for x in a) + " |\n")
        if args.appliquer:
            f.write("\nA faire ensuite : `python outils/rescanner_base.py` "
                    "(les @pseudos sont retroactifs sur toute la base).\n")

    print(f"{mode} : {len(retenues)} chaines, {n_comptes_nouveaux} comptes "
          f"nouveaux, {n_pseudos} pseudos au dictionnaire, "
          f"{len(inconnues)} inconnues")
    print(f"Ecrit : {chemin_md}")


if __name__ == "__main__":
    principal()
