# -*- coding: utf-8 -*-
"""Recolte les verdicts et les notes de Vincent depuis A_JUGER.xlsx.

C'est la moitie « retour » du critere 10 du contrat : chaque note laissee
dans la colonne « Ton commentaire » est relue, listee dans le rapport, et
accompagnee d'une ligne REPONSE que Claude complete en y donnant suite.
Une note sans reponse au rapport suivant = critere 10 echoue.

Les verdicts recoltes s'ajoutent a `donnees/jugements_recoltes.csv`
(cumulatif, une ligne par paire, le dernier verdict l'emporte). Ils
serviront de jeu de validation NEUF pour la regle R3.
"""

import csv
import datetime as dt
from pathlib import Path

import openpyxl

RACINE = Path(__file__).resolve().parent.parent
CHEMIN_CLASSEUR = RACINE / "A_JUGER.xlsx"
CHEMIN_RECOLTE = RACINE / "donnees" / "jugements_recoltes.csv"
DOSSIER_RECHERCHE = RACINE / "recherche"

CHAMPS = ["recolte_le", "video_id", "url", "chaine", "abonnes", "date",
          "entite", "signaux", "verdict", "commentaire"]


def lire_classeur():
    wb = openpyxl.load_workbook(CHEMIN_CLASSEUR, data_only=True)
    ws = wb["a juger"]
    lignes = list(ws.iter_rows(values_only=True))
    entetes = [str(c or "").strip() for c in lignes[0]]
    idx = {nom: entetes.index(nom) for nom in entetes if nom}
    recoltes = []
    for rang, l in enumerate(lignes[1:], start=2):
        def val(nom):
            i = idx.get(nom)
            return str(l[i]).strip() if i is not None and l[i] is not None else ""
        verdict, commentaire = val("TON VERDICT"), val("Ton commentaire")
        if not verdict and not commentaire:
            continue
        url = ""
        for cellule in ws[rang]:
            if cellule.hyperlink and cellule.hyperlink.target:
                url = cellule.hyperlink.target
                break
        video_id = url.split("watch?v=")[1].split("&")[0] if "watch?v=" in url else ""
        recoltes.append({
            "recolte_le": dt.date.today().isoformat(),
            "video_id": video_id,
            "url": url,
            "chaine": val("Chaine"),
            "abonnes": val("Abonnes"),
            "date": val("Date"),
            "entite": val("Entite possible"),
            "signaux": val("Signaux reperes"),
            "verdict": verdict,
            "commentaire": commentaire,
        })
    wb.close()
    return recoltes


def principal():
    recoltes = lire_classeur()

    existants = {}
    if CHEMIN_RECOLTE.exists():
        with open(CHEMIN_RECOLTE, encoding="utf-8-sig") as f:
            for l in csv.DictReader(f):
                existants[(l["video_id"], l["entite"])] = l
    nouveaux = 0
    for r in recoltes:
        cle = (r["video_id"], r["entite"])
        if cle not in existants or existants[cle]["verdict"] != r["verdict"] \
                or existants[cle]["commentaire"] != r["commentaire"]:
            nouveaux += 1
        existants[cle] = r
    with open(CHEMIN_RECOLTE, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=CHAMPS)
        w.writeheader()
        w.writerows(existants.values())

    aujourd_hui = dt.date.today().isoformat()
    chemin_rapport = DOSSIER_RECHERCHE / f"recolte_jugements_{aujourd_hui}.md"
    verdicts = {}
    for r in recoltes:
        if r["verdict"]:
            verdicts[r["verdict"]] = verdicts.get(r["verdict"], 0) + 1
    notes = [r for r in recoltes if r["commentaire"]]

    with open(chemin_rapport, "w", encoding="utf-8") as f:
        f.write(f"# Recolte des jugements — {aujourd_hui}\n\n")
        f.write("Produit par `outils/relire_jugements.py` depuis "
                "`A_JUGER.xlsx`.\n\n")
        f.write(f"**Verdicts recoltes : {sum(verdicts.values())}** "
                f"(lignes nouvelles ou modifiees : {nouveaux})\n\n")
        for v, n in sorted(verdicts.items(), key=lambda x: -x[1]):
            f.write(f"- {v} : {n}\n")
        f.write(f"\n## Notes de Vincent ({len(notes)}) — critere 10\n\n")
        if not notes:
            f.write("Aucune note ce passage.\n")
        for r in notes:
            f.write(f"**{r['chaine']} / {r['entite']}** ({r['video_id']}) — "
                    f"verdict : {r['verdict'] or '(aucun)'}\n")
            f.write(f"> {r['commentaire']}\n\n")
            f.write("REPONSE : (a completer par Claude en donnant suite)\n\n")
        f.write("\nCumul : `donnees/jugements_recoltes.csv` "
                f"({len(existants)} paires).\n")

    print(f"Verdicts recoltes : {sum(verdicts.values())}, "
          f"notes : {len(notes)}, nouveaux/modifies : {nouveaux}")
    print(f"Ecrit : {chemin_rapport}")


if __name__ == "__main__":
    principal()
