# -*- coding: utf-8 -*-
"""Le flux decouverte : ce que les lobbies disent eux-memes de leurs partenaires.

Principe (critere 11 du contrat) : quand un lobby publie une video, il nomme
souvent lui-meme ses partenaires — dans le titre, la description, ou en
taguant leur @compte. On ne DETECTE rien ici : on PROPOSE, preuve a l'appui,
et Vincent confirme.

Trois sources :
  1. les **@mentions** dans les videos de vitrines du corpus (identifiants
     stables — jamais de fragments de texte, lecon de JOURNAL 73) ;
  2. les **comptes connus** du registre (acquis/cartographie/COMPTES.xlsx,
     7 406 comptes), cherches par frontiere de mot dans ces memes videos ;
  3. la **moisson des 9 chaines de lobbies** faite par l'ancien projet le
     01/09 (343 lignes, createurs reconnus et identifiants) — on ne
     reconstruit pas ce qui existe.

Confiance : HAUTE pour un identifiant (@mention, pseudo, identifiant) ou un
nom en plusieurs mots ; BASSE pour un prenom ou mot seul — « Thomas » dans
un titre ne designe personne de sur. Le pre-tri est fourni, la decision
reste a Vincent.

Sorties : recherche/decouverte_vitrines_<date>.csv + .md, et le classeur
PROPOSITIONS.xlsx ou Vincent tranche (refuse d'ecraser des decisions non
recoltees).
"""

import csv
import datetime as dt
import re
import sqlite3
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

RACINE = Path(__file__).resolve().parent.parent
CHEMIN_BASE = RACINE / "donnees" / "veille.sqlite"
CHEMIN_VITRINES = RACINE / "donnees" / "contenus_vitrines.csv"
CHEMIN_COMPTES = RACINE / "acquis" / "cartographie" / "COMPTES.xlsx"
CHEMIN_DICO = RACINE / "DICTIONNAIRE.xlsx"
CHEMIN_MOISSON_LOBBIES = (RACINE / "acquis" / "donnees_brutes"
                          / "chaines_lobbies_2026-09-01.csv")
CHEMIN_PROPOSITIONS = RACINE / "PROPOSITIONS.xlsx"
DOSSIER_RECHERCHE = RACINE / "recherche"

MOTIF_MENTION = re.compile(r"@([A-Za-z0-9][A-Za-z0-9_.\-]{2,29})")
LONGUEUR_MIN_NOM = 5
DECISIONS = ["suivre ce compte", "ignorer", "je ne sais pas"]


def normaliser(texte):
    t = unicodedata.normalize("NFD", str(texte))
    t = "".join(c for c in t if unicodedata.category(c) not in ("Mn", "Cf"))
    return " ".join(t.lower().split())


def charger_comptes_connus():
    wb = openpyxl.load_workbook(CHEMIN_COMPTES, read_only=True, data_only=True)
    lignes = list(wb["comptes"].iter_rows(values_only=True))
    wb.close()
    entetes = [str(c or "").strip() for c in lignes[0]]
    i_plate = entetes.index("plateforme")
    i_ident = entetes.index("identifiant")
    i_pseudo = entetes.index("pseudo")
    i_nom = entetes.index("nom affiche")
    termes = {}
    for l in lignes[1:]:
        plateforme = str(l[i_plate] or "").strip()
        for source, type_terme in ((l[i_ident], "identifiant"),
                                   (l[i_pseudo], "pseudo"),
                                   (l[i_nom], "nom affiche")):
            terme = str(source or "").strip().lstrip("@")
            if len(terme) >= LONGUEUR_MIN_NOM:
                termes.setdefault(normaliser(terme),
                                  (terme, plateforme, type_terme))
    return termes


def exclusions_depuis_dictionnaire():
    """Comptes et noms des lobbies eux-memes : pas des propositions."""
    wb = openpyxl.load_workbook(CHEMIN_DICO, read_only=True, data_only=True)
    exclus = set()
    for nom_feuille, colonne in (("Signaux", "texte"), ("Entites", "entite")):
        lignes = list(wb[nom_feuille].iter_rows(values_only=True))
        entetes = [str(c or "").strip() for c in lignes[0]]
        i = entetes.index(colonne)
        for l in lignes[1:]:
            if l[i]:
                exclus.add(normaliser(str(l[i]).lstrip("@")))
    wb.close()
    return exclus


def confiance_pour(affichage, par_identifiant):
    if par_identifiant:
        return "haute"
    return "haute" if len(affichage.split()) >= 2 else "basse"


def decisions_existantes():
    if not CHEMIN_PROPOSITIONS.exists():
        return 0
    wb = openpyxl.load_workbook(CHEMIN_PROPOSITIONS, read_only=True,
                                data_only=True)
    n = 0
    if "propositions" in wb.sheetnames:
        lignes = list(wb["propositions"].iter_rows(values_only=True))
        if lignes:
            entetes = [str(c or "") for c in lignes[0]]
            if "TA DECISION" in entetes:
                i = entetes.index("TA DECISION")
                n = sum(1 for l in lignes[1:]
                        if i < len(l) and l[i] and str(l[i]).strip())
    wb.close()
    return n


def principal():
    n_dec = decisions_existantes()
    if n_dec and "--forcer" not in sys.argv:
        sys.exit(f"PROPOSITIONS.xlsx contient {n_dec} decisions non "
                 "recoltees : on ne l'ecrase pas.")

    exclus = exclusions_depuis_dictionnaire()
    connus = charger_comptes_connus()
    propositions = defaultdict(lambda: {
        "occurrences": 0, "videos": [], "entites": set(),
        "voies": set(), "affichage": "", "identifiant": ""})

    def proposer(cle, affichage, voie, entite, url, identifiant=""):
        cle = normaliser(cle)
        if not cle or cle in exclus:
            return
        p = propositions[cle]
        p["occurrences"] += 1
        p["affichage"] = p["affichage"] or affichage
        p["identifiant"] = p["identifiant"] or identifiant
        p["voies"].add(voie)
        if entite:
            p["entites"].add(entite)
        if url and len(p["videos"]) < 3 and url not in p["videos"]:
            p["videos"].append(url)

    # Sources 1 et 2 : les videos de vitrines du corpus.
    with open(CHEMIN_VITRINES, encoding="utf-8-sig") as f:
        vitrines = list(csv.DictReader(f))
    base = sqlite3.connect(CHEMIN_BASE)
    lues = 0
    motif_connus = re.compile(
        r"(?<!\w)(?:" + "|".join(
            re.escape(t) for t in sorted(connus, key=len, reverse=True))
        + r")(?!\w)")
    for v in vitrines:
        ligne = base.execute(
            "SELECT titre, description FROM videos WHERE video_id = ?",
            (v["video_id"],)).fetchone()
        if not ligne:
            continue
        lues += 1
        texte = f"{ligne[0] or ''}\n{ligne[1] or ''}"
        vus_ici = set()
        for m in MOTIF_MENTION.finditer(texte):
            pseudo = normaliser(m.group(1))
            if pseudo not in vus_ici:
                vus_ici.add(pseudo)
                proposer(pseudo, "@" + m.group(1), "@mention (vitrine)",
                         v["entite_du_canal"], v["url"],
                         identifiant=m.group(1))
        for m in motif_connus.finditer(normaliser(texte)):
            terme = m.group(0)
            if terme in vus_ici:
                continue
            vus_ici.add(terme)
            affichage, plateforme, type_terme = connus[terme]
            proposer(terme, affichage,
                     f"registre : {type_terme} {plateforme}",
                     v["entite_du_canal"], v["url"],
                     identifiant=(affichage if type_terme != "nom affiche"
                                  else ""))
    base.close()

    # Source 3 : la moisson des chaines de lobbies du 01/09.
    lignes_moisson = 0
    if CHEMIN_MOISSON_LOBBIES.exists():
        with open(CHEMIN_MOISSON_LOBBIES, encoding="utf-8-sig") as f:
            for r in csv.DictReader(f):
                lignes_moisson += 1
                noms = [n.strip() for n in
                        (r.get("createurs_nommes") or "").split("|")]
                idents = [i.strip() for i in
                          (r.get("identifiants") or "").split("|")]
                idents += [""] * (len(noms) - len(idents))
                for nom, ident in zip(noms, idents):
                    if not nom:
                        continue
                    proposer(ident or nom, nom,
                             "moisson chaines lobbies 01/09",
                             r.get("entite", ""), r.get("url", ""),
                             identifiant=ident)

    tri = sorted(propositions.items(), key=lambda x: -x[1]["occurrences"])
    aujourd_hui = dt.date.today().isoformat()

    chemin_csv = DOSSIER_RECHERCHE / f"decouverte_vitrines_{aujourd_hui}.csv"
    with open(chemin_csv, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["proposition", "identifiant", "confiance", "voies",
                    "occurrences", "entites", "preuves"])
        for _, p in tri:
            w.writerow([p["affichage"], p["identifiant"],
                        confiance_pour(p["affichage"], bool(p["identifiant"])),
                        " + ".join(sorted(p["voies"])), p["occurrences"],
                        " | ".join(sorted(p["entites"])),
                        " | ".join(p["videos"])])

    # Le classeur de decision de Vincent.
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "propositions"
    entetes = ["Proposition", "Identifiant", "Confiance", "Trouve par",
               "Occurrences", "Lobbies concernes", "Preuve 1", "Preuve 2",
               "TA DECISION", "Ton commentaire"]
    ws.append(entetes)
    for c in ws[1]:
        c.font = Font(bold=True)
    for _, p in tri:
        conf = confiance_pour(p["affichage"], bool(p["identifiant"]))
        ligne = [p["affichage"], p["identifiant"], conf,
                 " + ".join(sorted(p["voies"]))[:60], p["occurrences"],
                 " | ".join(sorted(p["entites"]))]
        ws.append(ligne + [None, None, None, None])
        rang = ws.max_row
        for i, url in enumerate(p["videos"][:2]):
            cellule = ws.cell(rang, 7 + i, "ouvrir")
            cellule.hyperlink = url
            cellule.font = Font(color="0563C1", underline="single")
    validation = DataValidation(type="list",
                                formula1='"' + ",".join(DECISIONS) + '"',
                                allow_blank=True, showDropDown=False)
    ws.add_data_validation(validation)
    validation.add(f"I2:I{ws.max_row}")
    for lettre, largeur in {"A": 26, "B": 20, "C": 10, "D": 40, "E": 12,
                            "F": 22, "G": 9, "H": 9, "I": 20, "J": 36}.items():
        ws.column_dimensions[lettre].width = largeur
    ws.freeze_panes = "A2"
    wb.save(CHEMIN_PROPOSITIONS)

    hautes = sum(1 for _, p in tri
                 if confiance_pour(p["affichage"], bool(p["identifiant"]))
                 == "haute")
    chemin_md = DOSSIER_RECHERCHE / f"decouverte_vitrines_{aujourd_hui}.md"
    with open(chemin_md, "w", encoding="utf-8") as f:
        f.write(f"# Flux decouverte — {aujourd_hui}\n\n")
        f.write("Produit par `outils/decouvrir_vitrines.py`.\n\n")
        f.write("| | |\n|---|---:|\n")
        f.write(f"| Videos de vitrines du corpus lues | {lues} |\n")
        f.write(f"| Lignes de la moisson lobbies du 01/09 | {lignes_moisson} |\n")
        f.write(f"| **Propositions distinctes** | **{len(tri)}** |\n")
        f.write(f"| dont confiance haute | {hautes} |\n")
        f.write(f"| dont confiance basse | {len(tri) - hautes} |\n\n")
        f.write("Vincent tranche dans `PROPOSITIONS.xlsx`. Rien n'entre en "
                "surveillance sans sa decision (critere 11). Detail : "
                f"`{chemin_csv.name}`.\n")

    print(f"Videos vitrines lues : {lues}  |  Moisson lobbies : "
          f"{lignes_moisson} lignes")
    print(f"Propositions : {len(tri)} ({hautes} haute confiance)")
    print(f"Ecrit : {CHEMIN_PROPOSITIONS}")
    print(f"Ecrit : {chemin_md}")


if __name__ == "__main__":
    principal()
